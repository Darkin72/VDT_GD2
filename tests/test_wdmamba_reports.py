"""Exercise WDMamba's real reporting pipeline without CUDA/Mamba weights."""

import contextlib
import csv
import io
import json
from pathlib import Path
import sys
import tempfile
import types
import unittest
from unittest.mock import patch

import cv2
import numpy as np
from openpyxl import load_workbook
import torch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "utils"))
import evaluate_wdmamba as evaluator


class WDMambaReportTests(unittest.TestCase):
    def test_four_workbooks_for_every_supported_dataset(self):
        cases = (
            ("i-haze", "I-HAZE", "ih01_hazy", "ih01", "real", "I-HAZE"),
            ("o-hazy", "O-HAZY", "41_outdoor", "41_outdoor", "real", "O-HAZE"),
            ("sots-indoor", "indoor", "001_0.8_1", "001", "synthetic", "SOTS-Indoor"),
            ("sots-outdoor", "outdoor", "001_0.8_0.08", "001", "synthetic", "SOTS-Outdoor"),
            ("nh-haze", None, "01_hazy", "01_GT", "real", "NH-HAZE"),
            ("dense-haze", None, "01_hazy", "01_GT", "real", "Dense-Haze"),
            ("haze4k", None, "001_a_b", "001", "synthetic", "Haze4K"),
            ("reside6k", None, "001", "001", "synthetic", "RESIDE-6K"),
            ("cdd11", None, "001", "001", "synthetic", "CDD-11"),
        )
        for dataset, folder, hazy_stem, clear_stem, origin, label in cases:
            with self.subTest(dataset=dataset), tempfile.TemporaryDirectory() as temporary:
                root = Path(temporary)
                data_root, output = root / "data", root / "results"
                checkpoint = root / "checkpoint.pth"
                checkpoint.touch()
                arguments = [
                    "evaluate_wdmamba.py", "--dataset", dataset,
                    "--data-root", str(data_root), "--checkpoint", str(checkpoint),
                    "--output-dir", str(output), "--device", "cpu", "--limit", "1",
                ]
                if dataset == "cdd11":
                    paired = data_root / "CDD-11_test"
                    clear_dir, hazy_dir = paired / "clear", paired / "low"
                    for category in evaluator.CDD11_CATEGORIES:
                        (paired / category).mkdir(parents=True)
                    arguments += ["--cdd11-test", str(paired)]
                elif dataset.startswith("sots-"):
                    paired = data_root / "Synthetic Objective Testing Set (SOTS) [RESIDE]" / folder / "test"
                    clear_dir, hazy_dir = paired / "clear", paired / "hazy"
                elif folder:
                    paired = data_root / folder / "test"
                    clear_dir, hazy_dir = paired / "clear", paired / "hazy"
                else:
                    paired = data_root / dataset / "test"
                    clear_dir, hazy_dir = paired / "GT", paired / "hazy"
                    arguments += ["--paired-root", str(paired)]
                clear_dir.mkdir(parents=True, exist_ok=True)
                hazy_dir.mkdir(parents=True, exist_ok=True)
                clear = np.random.default_rng(42).integers(30, 180, (32, 32, 3), dtype=np.uint8)
                hazy = clear + 10
                self.assertTrue(cv2.imwrite(str(clear_dir / (clear_stem + ".png")), clear))
                self.assertTrue(cv2.imwrite(str(hazy_dir / (hazy_stem + ".png")), hazy))
                # Only the checkpoint/model boundary is replaced. Image loading,
                # inference wrapper, metrics, CSV and all XLSX writers run normally.
                model = types.SimpleNamespace(restoration_network=torch.nn.Identity(), de_blocks=6)
                with patch.object(sys, "argv", arguments), patch.object(
                    evaluator, "load_model", return_value=model
                ), contextlib.redirect_stdout(io.StringIO()):
                    evaluator.main()

                expected = {"wdmamba_dataset.xlsx", "wdmamba_domain.xlsx", "wdmamba_fog.xlsx", "hardware.xlsx"}
                self.assertEqual({path.name for path in output.glob("*.xlsx")}, expected)
                with (output / "per_image.csv").open(encoding="utf-8-sig", newline="") as handle:
                    row = next(csv.DictReader(handle))
                self.assertEqual(row["data_origin"], origin)
                self.assertEqual(row["fog_level"], "light" if dataset.startswith("sots-") else "")
                expected_psnr = float(row["output_psnr"])

                for name, group in (("dataset", label), ("domain", origin.title()), ("fog", "Light fog")):
                    workbook = load_workbook(output / f"wdmamba_{name}.xlsx")
                    try:
                        sheet = workbook["Results"]
                        self.assertEqual(sheet["A4"].value, "WDMamba")
                        column = next(cell.column for cell in sheet[2] if cell.value == group)
                        value = sheet.cell(4, column).value
                        if name == "fog" and not dataset.startswith("sots-"):
                            self.assertIsNone(value)  # Do not invent a fog label.
                        else:
                            self.assertAlmostEqual(value, expected_psnr)
                        self.assertEqual(sheet.freeze_panes, "B4")
                    finally:
                        workbook.close()
                workbook = load_workbook(output / "hardware.xlsx")
                try:
                    sheet = workbook["Hardware"]
                    self.assertEqual(sheet["B1"].value, "WDMamba")
                    values = dict(sheet.iter_rows(min_row=2, values_only=True))
                    self.assertEqual(values["Images"], 1)
                    self.assertGreaterEqual(values["RAM peak (GB)"], 0)
                finally:
                    workbook.close()
                metadata = json.loads((output / "run_config.json").read_text(encoding="utf-8"))
                self.assertEqual({Path(path).name for path in metadata["excel_reports"].values()}, expected)


if __name__ == "__main__":
    unittest.main()
