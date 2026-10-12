"""The classic evaluators must satisfy the four-workbook benchmark contract."""

import json
from pathlib import Path
import subprocess
import sys

import cv2
import numpy as np
from openpyxl import load_workbook
import pytest


@pytest.mark.parametrize("solution", ["dcp", "cap"])
def test_classic_run_exports_hardware_and_preserves_performance(tmp_path, solution):
    data_root = tmp_path / "data"
    pair_root = data_root / "I-HAZE" / "test"
    image = np.random.default_rng(42).integers(30, 220, (32, 32, 3), dtype=np.uint8)
    for kind, name in (("clear", "ih1.png"), ("hazy", "ih1_hazy.png")):
        folder = pair_root / kind
        folder.mkdir(parents=True)
        assert cv2.imwrite(str(folder / name), image)
    root = Path(__file__).resolve().parents[1]
    subprocess.run([
        sys.executable, str(root / "utils" / "evaluate.py"),
        "--solution", solution, "--dataset", "i-haze", "--data-root", str(data_root),
        "--concurrent", "16", "--output-dir", str(tmp_path / "output"),
    ], check=True, capture_output=True, text=True, encoding="utf-8")
    output = tmp_path / "output" / solution / "test"
    assert {path.name for path in output.glob("*.xlsx")} == {
        f"{solution}_dataset.xlsx", f"{solution}_domain.xlsx",
        f"{solution}_fog.xlsx", "hardware.xlsx",
    }
    performance = json.loads((output / "performance.json").read_text(encoding="utf-8"))
    assert performance["images_completed"] == 1
    workbook = load_workbook(output / "hardware.xlsx", read_only=True)
    try:
        values = dict(workbook.active.iter_rows(values_only=True))
        assert values["Solution"] == solution.upper()
        assert values["images_completed"] == performance["images_completed"]
        assert values["fps"] == pytest.approx(performance["fps"])
        assert values["ram_peak_gb"] == performance["hardware"]["ram_peak_gb"]
    finally:
        workbook.close()
