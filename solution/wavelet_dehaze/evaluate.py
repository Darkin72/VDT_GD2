"""Evaluate HazeWaveNet and export reports, complexity and stage timings."""
from __future__ import annotations

import argparse
import json
import sys
import time
from pathlib import Path

import numpy as np
import torch
from PIL import Image
from skimage.metrics import peak_signal_noise_ratio, structural_similarity

PROJECT_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT_ROOT / "utils"))
from excel_report import export_excel_reports
from evaluation_hardware import write_hardware_report
from .model import HazeWaveNet

STAGE_KEYS = (
    "preprocess_ms", "dcp_fixed_prior", "dwt_mfde_decomposition",
    "low_frequency_fegg_fegb", "high_frequency_vp_vr", "wim_reconstruction",
    "final_refinement", "total", "total_with_preprocess",
)


def files(folder):
    return sorted(p for p in folder.iterdir() if p.is_file() and p.suffix.lower() in {".png", ".jpg", ".jpeg"})


def pairs(root: Path, dataset: str):
    base = root / "test"
    if dataset == "sots-indoor":
        base = root / "indoor" / "test"
    if dataset == "sots-outdoor":
        base = root / "outdoor" / "test"
    clear = {p.stem.lower(): p for p in files(base / "clear")}
    result = []
    for hazy in files(base / "hazy"):
        key = hazy.stem.lower()
        if dataset == "i-haze":
            key = key.removesuffix("_hazy")
        if dataset.startswith("sots"):
            key = key.split("_")[0]
        if key in clear:
            result.append((hazy, clear[key]))
    return result


def run_model(model, pair_list, device, max_side, profile_repeats):
    rows, elapsed_total = [], 0.0
    for index, (hazy_path, clear_path) in enumerate(pair_list, 1):
        started = time.perf_counter()
        with Image.open(hazy_path) as source:
            hazy_img = source.convert("RGB")
        with Image.open(clear_path) as source:
            clear_img = source.convert("RGB")
        if max_side and max(hazy_img.size) > max_side:
            scale = max_side / max(hazy_img.size)
            size = (round(hazy_img.width * scale), round(hazy_img.height * scale))
            hazy_img = hazy_img.resize(size, Image.Resampling.LANCZOS)
            clear_img = clear_img.resize(size, Image.Resampling.LANCZOS)
        hazy = torch.from_numpy(np.array(hazy_img, copy=True)).float().div(255).permute(2, 0, 1).unsqueeze(0).to(device)
        clear = np.array(clear_img, copy=True)
        preprocess_ms = (time.perf_counter() - started) * 1000.0
        if device == "cuda":
            torch.cuda.synchronize()
        output, stage = model.profile_forward(hazy, repeats=profile_repeats)
        if device == "cuda":
            torch.cuda.synchronize()
        elapsed_total += stage["total"] / 1000.0
        result = output[0].permute(1, 2, 0).mul(255).clamp(0, 255).byte().cpu().numpy()
        row = {
            "dataset": "",
            "image_index": index,
            "input_path": str(hazy_path),
            "target_path": str(clear_path),
            "output_psnr": peak_signal_noise_ratio(clear, result, data_range=255),
            "output_ssim": structural_similarity(clear, result, channel_axis=2, data_range=255),
            "runtime_ms": stage["total"],
            "preprocess_ms": preprocess_ms,
        }
        row.update(stage)
        row["preprocess_ms"] = preprocess_ms
        row["total_with_preprocess"] = preprocess_ms + stage["total"]
        print(f"    image {index}/{len(pair_list)} preprocess={preprocess_ms:.2f}ms inference={stage['total']:.2f}ms", flush=True)
        rows.append(row)
    return rows, elapsed_total


def model_complexity(model, device, input_side):
    from thop import profile
    parameters = sum(parameter.numel() for parameter in model.parameters())
    trainable = sum(parameter.numel() for parameter in model.parameters() if parameter.requires_grad)
    dummy = torch.zeros(1, 3, input_side, input_side, device=device)
    with torch.inference_mode():
        macs, _ = profile(model, inputs=(dummy,), verbose=False)
    return {
        "input_shape": f"1x3x{input_side}x{input_side}",
        "parameters": parameters,
        "trainable_parameters": trainable,
        "parameter_size_mb_fp32": parameters * 4 / 1024**2,
        "macs": macs,
        "gmacs": macs / 1e9,
        "flops": 2 * macs,
        "gflops": 2 * macs / 1e9,
    }


def write_detailed_workbook(path, rows, complexities, hardware):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    workbook = Workbook()
    per_image = workbook.active
    per_image.title = "PerImage"
    columns = ["dataset", "image_index", "input_path", "target_path", "output_psnr", "output_ssim", *STAGE_KEYS]
    per_image.append(columns)
    for row in rows:
        per_image.append([row.get(column) for column in columns])
    summary = workbook.create_sheet("DatasetSummary")
    summary_columns = ["dataset", "images", "psnr_mean", "ssim_mean", *STAGE_KEYS]
    summary.append(summary_columns)
    for dataset in sorted({row["dataset"] for row in rows}):
        subset = [row for row in rows if row["dataset"] == dataset]
        values = [float(np.mean([row[key] for row in subset])) for key in ("output_psnr", "output_ssim", *STAGE_KEYS)]
        summary.append([dataset, len(subset), *values])
    complexity_sheet = workbook.create_sheet("ModelComplexity")
    complexity_columns = ["dataset", "checkpoint", "input_shape", "parameters", "trainable_parameters", "parameter_size_mb_fp32", "macs", "gmacs", "flops", "gflops"]
    complexity_sheet.append(complexity_columns)
    for dataset, values in complexities.items():
        for value in values:
            complexity_sheet.append([dataset, value["checkpoint"], *[value[key] for key in complexity_columns[2:]]])
    hardware_sheet = workbook.create_sheet("Hardware")
    hardware_sheet.append(["key", "value"])
    for key, value in hardware.items():
        hardware_sheet.append([key, value])
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for column in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(column)].width = min(42, max(14, len(str(sheet.cell(1, column).value or "")) + 2))
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--data-root", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    for name in ("i-haze", "o-hazy", "sots-its", "sots-ots"):
        parser.add_argument(f"--checkpoint-{name}", type=Path, default=None)
    parser.add_argument("--device", choices=("cpu", "cuda"), default="cuda")
    parser.add_argument("--max-side", type=int, default=1024)
    parser.add_argument("--profile-repeats", type=int, default=10)
    args = parser.parse_args()
    if args.device == "cuda" and not torch.cuda.is_available():
        raise RuntimeError("CUDA is not available")
    if args.profile_repeats < 1:
        raise ValueError("--profile-repeats must be positive")
    names = {"i-haze": "I-HAZE", "o-hazy": "O-HAZY", "sots-indoor": "Synthetic Objective Testing Set (SOTS) [RESIDE]", "sots-outdoor": "Synthetic Objective Testing Set (SOTS) [RESIDE]"}
    checkpoints = {key: getattr(args, "checkpoint_" + key.replace("-", "_")) for key in ("i-haze", "o-hazy", "sots-its", "sots-ots")}
    if not any(checkpoints.values()):
        raise ValueError("At least one checkpoint argument is required")
    datasets = {"i-haze": "i-haze", "o-hazy": "o-hazy", "sots-its": "sots-indoor", "sots-ots": "sots-outdoor"}
    all_rows, total_time, total_images, complexities = [], 0.0, 0, {}
    for label, dataset in datasets.items():
        checkpoint = checkpoints[label]
        if checkpoint is None:
            continue
        print(f"[{label}] loading checkpoint={checkpoint}", flush=True)
        model = HazeWaveNet().to(args.device).eval()
        state = torch.load(checkpoint, map_location=args.device)
        model.load_state_dict(state.get("model", state))
        complexities[label] = [{"checkpoint": checkpoint.name, **model_complexity(model, args.device, side)} for side in (256, 1024)]
        root = args.data_root / ("I-HAZE" if dataset == "i-haze" else "O-HAZY" if dataset == "o-hazy" else names[dataset])
        pair_list = pairs(root, dataset)
        if not pair_list:
            raise ValueError(f"No matching test images for {dataset} under {root}")
        print(f"[{label}] images={len(pair_list)} profile_repeats={args.profile_repeats}", flush=True)
        rows, inference_time = run_model(model, pair_list, args.device, args.max_side, args.profile_repeats)
        for row, (hazy, _) in zip(rows, pair_list):
            stem = hazy.stem
            row["dataset"] = dataset
            row["data_origin"] = "synthetic" if dataset.startswith("sots") else "real"
            row["fog_level"] = "heavy" if "0.2" in stem else "medium" if any(x in stem for x in ("0.12", "0.16")) else "light"
        all_rows.extend(rows); total_time += inference_time; total_images += len(pair_list)
        timing_means = {key: float(np.mean([row[key] for row in rows])) for key in STAGE_KEYS}
        print(f"[{label}] mean_psnr={np.mean([row['output_psnr'] for row in rows]):.3f} mean_ssim={np.mean([row['output_ssim'] for row in rows]):.4f}", flush=True)
        print(f"[{label}] mean_timing_ms=" + ", ".join(f"{key}={value:.3f}" for key, value in timing_means.items()), flush=True)
        timing_path = args.output_dir / f"timing_{label.replace('-', '_')}.json"
        timing_path.parent.mkdir(parents=True, exist_ok=True)
        timing_path.write_text(json.dumps({"dataset": label, "images": len(rows), "mean_ms": timing_means}, indent=2), encoding="utf-8")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    paths = export_excel_reports("HazeWaveNet", all_rows, args.output_dir)
    performance = write_hardware_report(args.output_dir, "HazeWaveNet", total_images, total_time, total_time)
    write_detailed_workbook(args.output_dir / "hardware.xlsx", all_rows, complexities, performance["hardware"])
    print("Saved:"); [print(path) for path in paths.values()]
    print(args.output_dir / "hardware.xlsx")


if __name__ == "__main__":
    main()
